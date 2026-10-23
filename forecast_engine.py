"""Generic PhilRice Specific Tag forecasting engine (chronological 80/20 validation)."""
from __future__ import annotations
import math
import warnings
from pathlib import Path
import numpy as np
import pandas as pd
from statsmodels.tsa.arima.model import ARIMA
from statsmodels.tsa.holtwinters import Holt
from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.drawing.image import Image as ExcelImage
import io
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib.ticker import PercentFormatter, MultipleLocator
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

REQUIRED = ['Employee', 'Month', 'Specific Tag', '% of Total Workload']
CANONICAL_TAGS = ['core research function', 'core administration function', 'extension services', 'ad hoc', 'Research Enablement', 'RSS', 'SMART Farm', 'RCEF', 'RiceBIS', 'Malusog Rice', 'Special Project Support']
ORDERS = [(0,0,0),(1,0,0),(0,1,0),(0,0,1),(1,0,1),(1,1,0)]


def read_inputs(path: Path):
    """Read long-form monthly workload records from any Excel sheet with required headers."""
    book = pd.read_excel(path, sheet_name=None)
    frames = []
    for sheet, frame in book.items():
        frame.columns = [str(x).strip() for x in frame.columns]
        if set(REQUIRED).issubset(frame.columns):
            subset = frame[REQUIRED].copy()
            subset = subset.dropna(how='all')
            subset['source_sheet'] = sheet
            frames.append(subset)
    if not frames:
        raise ValueError('No sheet contains Employee, Month, Specific Tag, % of Total Workload. Use the included template.')
    df = pd.concat(frames, ignore_index=True)
    # Do not discard rows with blank workload shares silently: require an explicit 0%.
    missing_meta = df[['Employee','Month','Specific Tag']].isna().any(axis=1)
    if missing_meta.any():
        raise ValueError('Employee, Month and Specific Tag cannot be blank. Remove incomplete rows.')
    if df['% of Total Workload'].isna().any():
        raise ValueError('Blank workload % detected. Enter 0% for tags with no workload.')
    df['Employee'] = df['Employee'].astype(str).str.strip()
    df['Specific Tag'] = df['Specific Tag'].astype(str).str.strip()
    tag_lookup = {t.casefold(): t for t in CANONICAL_TAGS}
    df['Specific Tag'] = df['Specific Tag'].map(lambda t: tag_lookup.get(t.casefold(), t))
    dates = pd.to_datetime(df['Month'], errors='coerce')
    if dates.isna().any():
        raise ValueError('Some Month entries are invalid. Use dates such as January 2025.')
    df['Month'] = dates.dt.to_period('M').dt.to_timestamp()
    raw = df['% of Total Workload']
    text = raw.astype(str).str.strip()
    perc_text = text.str.endswith('%')
    val = pd.to_numeric(text.str.rstrip('%'), errors='coerce')
    if val.isna().any() or (val < 0).any():
        raise ValueError('Workload percentage must be a valid nonnegative number.')
    # Excel % cells are fractions; textual values like 35% are percentage points.
    fraction = val.copy()
    fraction[perc_text] /= 100
    # Ambiguous numeric values between 0 and 1 are treated as fractions.
    fraction[~perc_text & (val > 1)] /= 100
    if (fraction > 1.0001).any():
        raise ValueError('Workload percentages must not exceed 100%.')
    df['share'] = fraction.astype(float)
    result = {}
    for employee, group in df.groupby('Employee', sort=True):
        if group.duplicated(['Month','Specific Tag']).any():
            raise ValueError(f'{employee}: duplicate Month and Specific Tag entries; combine them before uploading.')
        wide = group.pivot(index='Month', columns='Specific Tag', values='share').sort_index()
        # Every monthly dataset includes ALL eleven PhilRice Specific Tags, even if
        # the source omitted their rows. Those missing tag/month pairs mean 0%.
        # Explicit zero workload values are preserved, never excluded from modeling.
        columns = CANONICAL_TAGS + sorted(set(wide.columns) - set(CANONICAL_TAGS))
        wide = wide.reindex(columns=columns).fillna(0.0)
        expected = pd.date_range(wide.index[0], wide.index[-1], freq='MS')
        if len(wide) < 6:
            raise ValueError(f'{employee}: at least 6 monthly observations are required; found {len(wide)}.')
        if len(expected) != len(wide):
            raise ValueError(f'{employee}: missing months in historical series. Include every month consecutively.')
        totals = wide.sum(axis=1)
        if ((totals - 1).abs() > .015).any():
            months = ', '.join(wide.index[(totals-1).abs() > .015].strftime('%b %Y')[:4])
            raise ValueError(f'{employee}: workload shares should add to 100% per month. Check {months}.')
        result[employee] = wide
    return result


def _fit_predict(series, kind, order=None):
    """Fit an ARIMA or Holt model to the observed history and predict one step."""
    with warnings.catch_warnings():
        warnings.simplefilter('ignore')
        if kind == 'ARIMA':
            fit = ARIMA(series, order=order, enforce_stationarity=False, enforce_invertibility=False).fit()
        else:
            fit = Holt(series, initialization_method='estimated').fit(optimized=True)
        return float(np.asarray(fit.forecast(1))[0])


def _future(series, kind, horizon, order=None):
    """Refit a chosen model to the full history and forecast a configurable horizon."""
    with warnings.catch_warnings():
        warnings.simplefilter('ignore')
        if kind == 'ARIMA':
            fit = ARIMA(series, order=order, enforce_stationarity=False, enforce_invertibility=False).fit()
        else:
            fit = Holt(series, initialization_method='estimated').fit(optimized=True)
        return np.asarray(fit.forecast(horizon), dtype=float)


def evaluate_series(series, horizon):
    """Compare ARIMA candidates against Holt via chronological expanding-window validation."""
    n = len(series)
    validation = max(1, math.ceil(n * .20))
    training = n - validation
    actual = np.asarray(series[training:], dtype=float)
    if np.allclose(series, 0):
        return dict(model='Constant zero', order=None, arima_rmse=0.0, holt_rmse=0.0,
                    selected_rmse=0.0, validation_actual=actual, validation_pred=np.zeros(validation),
                    forecast=np.zeros(horizon), train_months=training, validation_months=validation)
    if np.allclose(series, series[0]):
        return dict(model='Constant', order=None, arima_rmse=0.0, holt_rmse=0.0,
                    selected_rmse=0.0, validation_actual=actual, validation_pred=np.repeat(series[0],validation),
                    forecast=np.repeat(series[0],horizon), train_months=training, validation_months=validation)
    candidates = [('ARIMA', order) for order in ORDERS if training >= max(4, sum(order)+2)] + [('Holt',None)]
    best = None
    arima_best = float('inf'); holt_score = float('inf')
    for kind, order in candidates:
        preds = []
        for at in range(training, n):
            history = np.asarray(series[:at], dtype=float)
            try:
                prediction = _fit_predict(history, kind, order)
                if not np.isfinite(prediction):
                    raise ValueError('Nonfinite prediction')
            except Exception:
                preds = []; break
            preds.append(prediction)
        if len(preds) != validation:
            continue
        score = float(np.sqrt(np.mean((actual - np.asarray(preds))**2)))
        if kind == 'ARIMA': arima_best = min(arima_best, score)
        else: holt_score = score
        if best is None or score < best[0]: best = (score, kind, order, np.asarray(preds))
    if best is None:
        raise ValueError('No ARIMA or Holt model could be fitted to this series. Check the input history.')
    score, kind, order, preds = best
    try:
        forecast = _future(np.asarray(series, dtype=float), kind, horizon, order)
    except Exception:
        forecast = np.repeat(series[-1], horizon)
    return dict(model=kind, order=order, arima_rmse=arima_best, holt_rmse=holt_score,
                selected_rmse=score, validation_actual=actual, validation_pred=preds,
                forecast=np.maximum(0,forecast), train_months=training, validation_months=validation)


def run_all(datasets, horizon, report):
    """Forecast each employee/tag, updating the supplied progress callback after every tag."""
    jobs = sum(len(m.columns) for m in datasets.values())
    done = 0; result = {}
    for employee, matrix in datasets.items():
        tags = {}
        for tag in matrix.columns:
            tags[tag] = evaluate_series(matrix[tag].to_numpy(float), horizon)
            done += 1
            report(done, jobs, f'{employee} — {tag} ({done}/{jobs})')
        dates = pd.date_range(matrix.index[-1] + pd.offsets.MonthBegin(1), periods=horizon, freq='MS')
        raw = np.column_stack([tags[t]['forecast'] for t in matrix.columns])
        total = raw.sum(axis=1)
        for i in range(horizon):
            if total[i] > 0: raw[i,:] /= total[i]
            else: raw[i,:] = matrix.iloc[-1].to_numpy(float)
        for j,tag in enumerate(matrix.columns):
            tags[tag]['forecast_normalized'] = raw[:,j]
        result[employee] = {'history':matrix, 'tags':tags, 'dates':dates}
    return result



def make_chart_image(dates, lines, title, every=1, width=11.8, height=4.9):
    """Render a legible Excel-compatible chart with explicit date and percent ticks."""
    fig, ax = plt.subplots(figsize=(width, height), dpi=130)
    x = np.arange(len(dates))
    for name, values in lines.items():
        ax.plot(x, np.asarray(values, dtype=float), linewidth=1.7, marker='o' if len(x)<13 else None,
                markersize=3, label=name)
    ax.set_xlim(-0.3, max(1,len(x)-0.7))
    ax.set_ylim(0, 1.02)
    ax.set_title(title, fontsize=12, weight='bold', pad=12)
    ax.set_ylabel('% of Total Workload', fontsize=10, labelpad=9)
    ax.set_xlabel('Month and Year', fontsize=10, labelpad=9)
    ax.yaxis.set_major_locator(MultipleLocator(0.1))
    ax.yaxis.set_major_formatter(PercentFormatter(xmax=1.0, decimals=0))
    tick_idx = np.arange(0, len(x), max(1, every))
    if len(x) and len(x)-1 not in tick_idx: tick_idx = np.append(tick_idx,len(x)-1)
    labels = [pd.Timestamp(d).strftime('%b %Y') for d in dates]
    ax.set_xticks(tick_idx, [labels[i] for i in tick_idx], rotation=35, ha='right', fontsize=9)
    ax.tick_params(axis='y', labelsize=9, labelleft=True)
    ax.tick_params(axis='x', labelbottom=True, bottom=True)
    ax.grid(axis='y', alpha=.22)
    ax.spines['top'].set_visible(False)
    ax.spines['right'].set_visible(False)
    ax.legend(loc='upper center', bbox_to_anchor=(.5,-.31), ncol=3,
              frameon=False, fontsize=7.5)
    fig.subplots_adjust(left=.105, right=.97, top=.87, bottom=.29)
    stream = io.BytesIO()
    fig.savefig(stream, format='png', dpi=130, facecolor='white')
    plt.close(fig)
    stream.seek(0)
    img = ExcelImage(stream)
    img._philrice_buffer = stream
    return img

def write_workbook(results, path):
    """Create client-facing Excel forecast, model-validation tables, and labeled charts."""
    wb = Workbook(); overview = wb.active; overview.title = 'Overview'
    overview.append(['PhilRice Workload Forecast — Model Validation and Future Shares'])
    overview.append(['Employee','Specific Tag','History Months','Train Months','Validation Months','Model','ARIMA RMSE','Holt RMSE','Selected RMSE'])
    for employee, item in results.items():
        matrix = item['history']; dates = item['dates']; tags = item['tags']
        for tag, fit in tags.items():
            overview.append([employee,tag,len(matrix),fit['train_months'],fit['validation_months'],
                             fit['model']+(f' {fit["order"]}' if fit['order'] else ''),
                             None if np.isinf(fit['arima_rmse']) else fit['arima_rmse'],
                             None if np.isinf(fit['holt_rmse']) else fit['holt_rmse'],fit['selected_rmse']])
        safe = ''.join(c if c not in '[]:*?/\\' else '_' for c in employee)[:23]
        sh = wb.create_sheet(safe + ' Forecast')
        sh.append([f'Forecast — {employee}']); sh.append(['Month'] + list(matrix.columns))
        for i,d in enumerate(dates):
            sh.append([d.strftime('%b %Y')] + [float(tags[t]['forecast_normalized'][i]) for t in matrix.columns])
            # Month/year is stored as a readable chart category label.
            for c in range(2,sh.max_column+1): sh.cell(sh.max_row,c).number_format='0.0%'
        forecast_lines = {tag: tags[tag]['forecast_normalized'] for tag in matrix.columns}
        interval = max(1, math.ceil(len(dates)/12))
        image = make_chart_image(dates,forecast_lines,
                                 f'{employee} — Monthly Specific Tag Forecast', every=interval)
        sh.add_image(image, f'A{sh.max_row+4}')
        val=wb.create_sheet(safe + ' Validation'); val.append([f'Historical actual vs rolling forecast — {employee}'])
        val.append(['Specific Tag','Month','Actual','Validation Forecast','Selected Model','RMSE'])
        for tag, fit in tags.items():
            vd=matrix.index[-fit['validation_months']:]
            for i,d in enumerate(vd):
                val.append([tag,d.to_pydatetime(),float(fit['validation_actual'][i]),float(fit['validation_pred'][i]),fit['model'],fit['selected_rmse']])
                val.cell(val.max_row,2).number_format='mmmm yyyy'
                val.cell(val.max_row,3).number_format='0.0%';val.cell(val.max_row,4).number_format='0.0%'
        start=val.max_row+3
        for tag, fit in tags.items():
            val.cell(start,1, f'Validation Chart — {tag} (including zero-series)')
            val.cell(start+1,1,'Month'); val.cell(start+1,2,'Actual');val.cell(start+1,3,'Forecast')
            vd=matrix.index[-fit['validation_months']:]
            for i,d in enumerate(vd):
                val.cell(start+2+i,1,d.strftime('%b %Y'))
                val.cell(start+2+i,2,float(fit['validation_actual'][i])).number_format='0.0%'
                val.cell(start+2+i,3,float(fit['validation_pred'][i])).number_format='0.0%'
            image = make_chart_image(vd, {
                'Actual':fit['validation_actual'], 'Forecast':fit['validation_pred']},
                f'{employee} — {tag} — Validation', every=max(1, math.ceil(len(vd)/10)),
                width=9.5, height=4.5)
            val.add_image(image, f'H{start}')
            start += max(32,len(vd)+6)
    for sh in wb:
        sh.freeze_panes='B3'
        for col in sh.columns:
            letter=get_column_letter(col[0].column)
            sh.column_dimensions[letter].width=min(48,max(15,max(len(str(c.value or '')) for c in list(col)[:min(150,len(col))])+2))
        for cell in sh[2]:
            cell.fill=PatternFill('solid', fgColor='244D71');cell.font=Font(color='FFFFFF',bold=True)
    wb.save(path)
